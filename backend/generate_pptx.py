import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path):
    prs = Presentation()
    
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Media paths
    media_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "extracted_ppt_images", "ppt", "media")
    
    def get_media_path(filename):
        path = os.path.join(media_dir, filename)
        return path if os.path.exists(path) else None

    # ════════════════════════════════════════════════
    # CLEAN LIGHT ACADEMIC PALETTE (Matches PDF)
    # ════════════════════════════════════════════════
    BG_LIGHT = RGBColor(248, 250, 252)       # Light Gray-Blue Background (#F8FAFC)
    CARD_BG = RGBColor(255, 255, 255)        # Pure White Card (#FFFFFF)
    CARD_BG_ALT = RGBColor(241, 245, 249)    # Slight Gray Card (#F1F5F9)
    
    # Accent Colors from PDF
    ACCENT_BLUE = RGBColor(59, 130, 246)     # Primary Blue
    ACCENT_DARK_BLUE = RGBColor(30, 58, 138) # Dark Blue
    ACCENT_PINK = RGBColor(219, 39, 119)     # Pink/Magenta
    ACCENT_GREEN = RGBColor(16, 185, 129)    # Emerald Green
    ACCENT_ORANGE = RGBColor(245, 158, 11)   # Orange/Warning
    BORDER_MUTED = RGBColor(226, 232, 240)   # Light Border
    
    TEXT_DARK = RGBColor(15, 23, 42)         # Very Dark Slate (Headings)
    TEXT_MUTED = RGBColor(71, 85, 105)       # Slate Gray (Body)
    TEXT_BLUE = RGBColor(37, 99, 235)        # Accent Blue Text
    
    FONT_TITLE = "Calibri"
    FONT_BODY = "Calibri"
    FONT_MONO = "Consolas"

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_LIGHT

    def add_slide_header(slide, title_text):
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(32)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK

    def add_card(slide, left, top, width, height, title="", fill_color=CARD_BG, border_color=BORDER_MUTED, title_color=TEXT_DARK, icon=""):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.18), width - Inches(0.4), Inches(0.5))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            if icon:
                p.text = f"{icon} {title}"
            else:
                p.text = title
            p.font.name = FONT_TITLE
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = title_color
            return left + Inches(0.2), top + Inches(0.75), width - Inches(0.4), height - Inches(0.9)
            
        return left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4)

    def add_stat_box(slide, left, top, width, height, stat_number, stat_label, subtext="", border_color=ACCENT_BLUE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        
        tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.15), width - Inches(0.3), height - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        
        p1 = tf.paragraphs[0]
        p1.text = str(stat_number)
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(36)
        p1.font.bold = True
        p1.font.color.rgb = border_color
        p1.space_after = Pt(2)
        
        p2 = tf.add_paragraph()
        p2.text = stat_label
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK
        
        if subtext:
            p3 = tf.add_paragraph()
            p3.text = subtext
            p3.font.name = FONT_BODY
            p3.font.size = Pt(11)
            p3.font.color.rgb = TEXT_MUTED

    slide_layout = prs.slide_layouts[6]

    # ════════════════════════════════════════════════
    # SLIDE 1: Title Slide (Matches PDF Page 1)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    
    img_student = get_media_path("image-1-1.jpeg")
    if img_student:
        slide.shapes.add_picture(img_student, Inches(0.8), Inches(1.5), width=Inches(3.2), height=Inches(4.27))
    
    tb_title = slide.shapes.add_textbox(Inches(4.5), Inches(1.2), Inches(8.0), Inches(2.0))
    tf = tb_title.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "EvalTutor – AI-Powered\nMultilingual Tutoring Evaluation\nPlatform"
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_DARK
    p1.space_after = Pt(14)
    
    p2 = tf.add_paragraph()
    p2.text = "B.Tech Final Year Project Presentation | Academic Year 2025–2026"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(16)
    p2.font.color.rgb = TEXT_MUTED

    m1_l, m1_t, m1_w, m1_h = add_card(slide, Inches(4.5), Inches(4.3), Inches(3.8), Inches(2.2), "Student Details", ACCENT_BLUE, ACCENT_BLUE, RGBColor(255,255,255))
    tb_m1 = slide.shapes.add_textbox(m1_l, m1_t, m1_w, m1_h)
    tf_m1 = tb_m1.text_frame
    tf_m1.word_wrap = True
    
    student_info = [
        ("Student Name", "[MOHAN T]"),
        ("Register Number", "[710024243701]"),
        ("Department", "[B.tech- AI & DS]")
    ]
    for idx, (label, val) in enumerate(student_info):
        p = tf_m1.add_paragraph() if idx > 0 else tf_m1.paragraphs[0]
        p.text = f"• {label}: "
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(255,255,255)
        run = p.add_run()
        run.text = val
        run.font.bold = True
        run.font.color.rgb = RGBColor(255,255,255)
        p.space_after = Pt(6)

    m2_l, m2_t, m2_w, m2_h = add_card(slide, Inches(8.5), Inches(4.3), Inches(4.0), Inches(2.2), "Institutional Details", CARD_BG, BORDER_MUTED, TEXT_DARK)
    tb_m2 = slide.shapes.add_textbox(m2_l, m2_t, m2_w, m2_h)
    tf_m2 = tb_m2.text_frame
    tf_m2.word_wrap = True
    
    inst_info = [
        ("College Name", "[ANNA UNIVERSITY REGIONAL CAMPUS]"),
        ("Guide", "[Dr.Marikannu]"),
        ("Academic Year", "2025–2026")
    ]
    for idx, (label, val) in enumerate(inst_info):
        p = tf_m2.add_paragraph() if idx > 0 else tf_m2.paragraphs[0]
        p.text = f"• {label}: "
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        run = p.add_run()
        run.text = val
        run.font.bold = True
        run.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # ════════════════════════════════════════════════
    # SLIDE 2: The Problem (Matches PDF Page 2)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "The Problem: AI Tutors Are Unmonitored & Unreliable")

    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "AI tutors now serve millions of students globally — yet there is zero pedagogical oversight on the content they generate. The consequences are serious and growing."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_MUTED

    problems = [
        ("🚀 Explosive AI Adoption", "AI tutors serve millions of students globally, yet response quality remains entirely unchecked and unmonitored.", ACCENT_BLUE),
        ("🧠 Hallucinations Are Rampant", "LLMs confidently generate incorrect facts, misleading students during critical learning phases with no warning.", ACCENT_BLUE),
        ("⚠️ Dangerous Dependency", "Students copy AI answers without understanding — directly undermining conceptual learning and critical thinking.", ACCENT_PINK),
        ("🌐 Language Barrier", "Most AI tutors operate only in English, excluding 1.5 billion+ regional-language learners worldwide.", ACCENT_BLUE),
        ("🔍 No Evaluation Layer", "There is currently zero pedagogical oversight or quality scoring on AI-generated tutoring content.", ACCENT_BLUE),
        ("✅ EvalTutor Fills the Gap", "A real-time evaluation platform that monitors, scores, and continuously improves AI tutoring quality.", ACCENT_PINK)
    ]

    for idx, (p_title, p_desc, p_color) in enumerate(problems):
        col = idx % 2
        row = idx // 2
        c_l = Inches(0.8 + col * 5.95)
        c_t = Inches(2.0 + row * 1.7)
        
        cl, ct, cw, ch = add_card(slide, c_l, c_t, Inches(5.7), Inches(1.5), p_title, CARD_BG, p_color, TEXT_DARK)
        tb = slide.shapes.add_textbox(cl, ct, cw, ch)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = p_desc
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED

    # ════════════════════════════════════════════════
    # SLIDE 3: Objectives (Matches PDF Page 3)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "Objectives")

    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "EvalTutor is designed around five core objectives that together create a comprehensive, trustworthy AI tutoring ecosystem."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_MUTED

    pillars = [
        ("Monitor AI Quality", "Score every AI response for pedagogical accuracy, clarity, and educational value on a 0–100 scale.", ACCENT_BLUE, "🛡️"),
        ("Detect Hallucinations", "Flag responses with high hallucination risk before students act on potentially incorrect information.", ACCENT_BLUE, "🚩"),
        ("Measure Dependency", "Track over-reliance on AI answers versus independent reasoning to promote genuine learning.", ACCENT_PINK, "🎯"),
        ("Multilingual Support", "Enable seamless tutoring in regional languages via automatic language detection and translation.", ACCENT_BLUE, "🌐"),
        ("Analytics Dashboard", "Provide students and educators with real-time learning insights, progress metrics, and weak topic identification.", ACCENT_BLUE, "📈")
    ]

    for idx, (title, desc, color, icon) in enumerate(pillars):
        if idx < 3:
            col = idx
            row = 0
            c_l = Inches(0.8 + col * 3.9)
            c_t = Inches(2.0)
            c_w = Inches(3.7)
            c_h = Inches(2.5)
        else:
            col = idx - 3
            row = 1
            c_l = Inches(0.8 + col * 5.95)
            c_t = Inches(4.7)
            c_w = Inches(5.7)
            c_h = Inches(2.5)
            
        cl, ct, cw, ch = add_card(slide, c_l, c_t, c_w, c_h, title, CARD_BG, color, TEXT_DARK, icon=icon)
        
        tb = slide.shapes.add_textbox(cl, ct, cw, ch)
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_desc = tf.paragraphs[0]
        p_desc.text = desc
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = TEXT_MUTED

    # ════════════════════════════════════════════════
    # SLIDE 4: Workflow (Matches PDF Page 4)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "How EvalTutor Works: End-to-End Workflow")

    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "A seamless 7-step pipeline transforms a student's question — in any language — into an evaluated, scored, and translated AI tutoring response delivered in real time."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_MUTED

    img_wf = get_media_path("image-4-1.png")
    if img_wf:
        slide.shapes.add_picture(img_wf, Inches(0.95), Inches(2.0), width=Inches(11.433), height=Inches(4.0))

    tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.6))
    tf_foot = tb_foot.text_frame
    tf_foot.word_wrap = True
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "Each step is automated and executes in milliseconds — the student experiences a seamless, real-time tutoring interaction while the evaluation engine works silently in the background to score every response."
    p_foot.font.name = FONT_BODY
    p_foot.font.size = Pt(12)
    p_foot.font.color.rgb = TEXT_MUTED

    # ════════════════════════════════════════════════
    # SLIDE 5: System Architecture (Matches PDF Page 5)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "System Architecture")

    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "EvalTutor's architecture is modular and layered — each component is independently scalable and communicates through a FastAPI orchestration backbone."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_MUTED

    img_arch = get_media_path("image-5-1.jpeg")
    if img_arch:
        slide.shapes.add_picture(img_arch, Inches(0.95), Inches(2.0), width=Inches(11.433), height=Inches(3.3))

    m1_l, m1_t, m1_w, m1_h = add_card(slide, Inches(0.8), Inches(5.5), Inches(5.7), Inches(1.6), "Core Processing Modules", ACCENT_BLUE, ACCENT_BLUE, RGBColor(255,255,255))
    tb_m1 = slide.shapes.add_textbox(m1_l, m1_t, m1_w, m1_h)
    tf_m1 = tb_m1.text_frame
    tf_m1.word_wrap = True
    
    core_mods = [
        ("Langdetect", "Identifies input language in milliseconds"),
        ("Googletrans", "Bridges regional languages ↔ English"),
        ("Groq Llama-3.3-70B", "Generates pedagogically structured responses")
    ]
    for idx, (label, val) in enumerate(core_mods):
        p = tf_m1.add_paragraph() if idx > 0 else tf_m1.paragraphs[0]
        p.text = f"• "
        run_l = p.add_run()
        run_l.text = f"{label}: "
        run_l.font.bold = True
        run_l.font.color.rgb = RGBColor(255,255,255)
        run_v = p.add_run()
        run_v.text = val
        run_v.font.bold = False
        run_v.font.color.rgb = RGBColor(255,255,255)
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.space_after = Pt(4)

    m2_l, m2_t, m2_w, m2_h = add_card(slide, Inches(6.8), Inches(5.5), Inches(5.7), Inches(1.6), "Evaluation & Output Layer", CARD_BG, BORDER_MUTED, TEXT_DARK)
    tb_m2 = slide.shapes.add_textbox(m2_l, m2_t, m2_w, m2_h)
    tf_m2 = tb_m2.text_frame
    tf_m2.word_wrap = True
    
    eval_mods = [
        ("Evaluation Engine", "Scores quality, detects hallucinations, measures dependency"),
        ("Dashboard", "Chart.js-powered real-time analytics"),
        ("FastAPI", "Handles all routing and session management")
    ]
    for idx, (label, val) in enumerate(eval_mods):
        p = tf_m2.add_paragraph() if idx > 0 else tf_m2.paragraphs[0]
        p.text = f"• "
        run_l = p.add_run()
        run_l.text = f"{label}: "
        run_l.font.bold = True
        run_l.font.color.rgb = TEXT_DARK
        run_v = p.add_run()
        run_v.text = val
        run_v.font.bold = False
        run_v.font.color.rgb = TEXT_MUTED
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.space_after = Pt(4)

    # ════════════════════════════════════════════════
    # SLIDE 6: Technologies Used (Matches PDF Page 6)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "Technologies Used")

    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "EvalTutor is built on a modern, lightweight, and high-performance technology stack — each layer chosen for speed, accuracy, and developer productivity."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_MUTED

    t_l, t_t, t_w, t_h = add_card(slide, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8), "", CARD_BG, BORDER_MUTED, TEXT_DARK)
    tb_t = slide.shapes.add_textbox(t_l, t_t, t_w, t_h)
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    
    # Table header styling
    p_h = tf_t.paragraphs[0]
    p_h.text = "Layer                                 Technology"
    p_h.font.name = FONT_TITLE
    p_h.font.size = Pt(16)
    p_h.font.bold = True
    p_h.font.color.rgb = ACCENT_BLUE
    p_h.space_after = Pt(14)

    tech_items = [
        ("Frontend", "HTML, CSS, JavaScript"),
        ("Backend", "FastAPI, Python"),
        ("AI Model", "Groq Llama-3.3-70B-Versatile"),
        ("Translation", "Googletrans, Langdetect"),
        ("Charts", "Chart.js"),
        ("Prompt Engineering", "Structured JSON & Markdown")
    ]
    for layer, tech in tech_items:
        p = tf_t.add_paragraph()
        run_l = p.add_run()
        run_l.text = f"{layer:<25}"
        run_l.font.color.rgb = TEXT_DARK
        run_l.font.bold = True
        
        run_t = p.add_run()
        run_t.text = tech
        run_t.font.color.rgb = TEXT_MUTED
        
        p.font.name = FONT_BODY
        p.font.size = Pt(14)
        p.space_after = Pt(14)

    c1_left, c1_top, c1_width, c1_height = add_card(slide, Inches(6.8), Inches(2.0), Inches(5.7), Inches(1.4), "⚡ Ultra-Fast Inference", CARD_BG, BORDER_MUTED, TEXT_DARK)
    tb_c1 = slide.shapes.add_textbox(c1_left, c1_top, c1_width, c1_height)
    tf_c1 = tb_c1.text_frame
    p_c1 = tf_c1.paragraphs[0]
    p_c1.text = "Groq's hardware-accelerated API delivers AI responses in milliseconds — enabling truly real-time tutoring interactions."
    p_c1.font.name = FONT_BODY
    p_c1.font.size = Pt(12)
    p_c1.font.color.rgb = TEXT_MUTED

    c2_left, c2_top, c2_width, c2_height = add_card(slide, Inches(6.8), Inches(3.6), Inches(5.7), Inches(1.4), "🐍 Python-First Stack", CARD_BG, BORDER_MUTED, TEXT_DARK)
    tb_c2 = slide.shapes.add_textbox(c2_left, c2_top, c2_width, c2_height)
    tf_c2 = tb_c2.text_frame
    p_c2 = tf_c2.paragraphs[0]
    p_c2.text = "FastAPI and Python provide a clean, async-ready backend that integrates seamlessly with all AI and translation modules."
    p_c2.font.name = FONT_BODY
    p_c2.font.size = Pt(12)
    p_c2.font.color.rgb = TEXT_MUTED
    
    c3_left, c3_top, c3_width, c3_height = add_card(slide, Inches(6.8), Inches(5.2), Inches(5.7), Inches(1.4), "📊 Live Visualization", CARD_BG, BORDER_MUTED, TEXT_DARK)
    tb_c3 = slide.shapes.add_textbox(c3_left, c3_top, c3_width, c3_height)
    tf_c3 = tb_c3.text_frame
    p_c3 = tf_c3.paragraphs[0]
    p_c3.text = "Chart.js renders dynamic, real-time analytics charts on the student dashboard after every single interaction."
    p_c3.font.name = FONT_BODY
    p_c3.font.size = Pt(12)
    p_c3.font.color.rgb = TEXT_MUTED

    # ════════════════════════════════════════════════
    # SLIDE 7: Key Features (Matches PDF Page 7)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "Key Features")

    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Seven tightly integrated features make EvalTutor a complete, end-to-end AI tutoring quality platform — not just a chatbot wrapper."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_MUTED

    features = [
        ("⚡ Live AI Chat", "Real-time tutoring powered by Groq's ultra-fast inference API with structured educational responses.", ACCENT_BLUE),
        ("🌐 Multilingual Support", "Auto-detects and responds in the student's native language — inclusive for 1.5B+ regional learners.", ACCENT_BLUE),
        ("📊 Pedagogical Evaluation", "Every response scored against educational quality standards on a 0–100 scale in real time.", ACCENT_PINK),
        ("🚨 Hallucination Detection", "Flags factually risky AI responses (Low / Medium / High) before they mislead students.", ACCENT_BLUE),
        ("🔗 Dependency Tracking", "Monitors how often students rely on AI versus thinking independently — promoting genuine learning.", ACCENT_BLUE),
        ("📈 Student Dashboard", "Personalized view of quality scores, weak topics, and learning progress over time.", ACCENT_PINK)
    ]

    for idx, (title, desc, color) in enumerate(features):
        col = idx % 3
        row = idx // 3
        c_l = Inches(0.8 + col * 3.97)
        c_t = Inches(2.0 + row * 2.2)
        
        cl, ct, cw, ch = add_card(slide, c_l, c_t, Inches(3.77), Inches(2.0), title, CARD_BG, color, TEXT_DARK)
        tb = slide.shapes.add_textbox(cl, ct, cw, ch)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED

    c7_left, c7_top, c7_width, c7_height = add_card(slide, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.8), "🔄 Real-Time Analytics", CARD_BG, ACCENT_BLUE, TEXT_DARK)
    tb_c7 = slide.shapes.add_textbox(c7_left, c7_top, c7_width, c7_height)
    tf_c7 = tb_c7.text_frame
    p_c7 = tf_c7.paragraphs[0]
    p_c7.text = "Live charts and metrics updated after every single interaction using Chart.js visualizations."
    p_c7.font.name = FONT_BODY
    p_c7.font.size = Pt(12)
    p_c7.font.color.rgb = TEXT_MUTED

    # ════════════════════════════════════════════════
    # SLIDE 8: Working Methodology (Matches PDF Page 8)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "Working Methodology: 10-Step Process")
    
    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "From question submission to dashboard update — EvalTutor's complete processing pipeline executes automatically in under two seconds per interaction."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_MUTED

    steps_10 = [
        ("01", "Student Submits Question", "Student types a question in any language via the chat interface."),
        ("02", "Language Identification", "Langdetect identifies the input language automatically."),
        ("03", "Query Translation", "Googletrans translates the query to English for AI processing."),
        ("04", "AI Response Generation", "Groq Llama-3.3 generates a structured, pedagogically sound AI response."),
        ("05", "Response Translation Back", "Response is translated back to the student's native language."),
        ("06", "Evaluation Engine Analysis", "Evaluation Engine analyzes response quality across multiple dimensions."),
        ("07", "Quality Score Calculated", "Quality Score is computed on a 0–100 scale using a multi-dimensional rubric."),
        ("08", "Hallucination Risk Flagged", "Hallucination Risk is classified as Low, Medium, or High per response."),
        ("09", "Dependency Level Logged", "Dependency Level is measured and recorded for the session."),
        ("10", "Dashboard Updated", "All metrics are pushed to the student dashboard in real time via Chart.js.")
    ]

    for idx, (num, title, desc) in enumerate(steps_10):
        col = idx % 2
        row = idx // 2
        c_l = Inches(0.8 + col * 5.95)
        c_t = Inches(2.0 + row * 1.05)
        
        # Simple line separator style like PDF
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_l, c_t, Inches(5.7), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = ACCENT_BLUE if idx % 3 != 0 else ACCENT_PINK
        line.line.fill.background()
        
        tb = slide.shapes.add_textbox(c_l, c_t + Inches(0.05), Inches(5.7), Inches(0.9))
        tf = tb.text_frame
        
        p1 = tf.paragraphs[0]
        p1.text = f"{num}  {title}"
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p1.space_after = Pt(4)
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED

    # ════════════════════════════════════════════════
    # SLIDE 9: Results & Dashboard Insights (Matches PDF Page 9)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "Results & Dashboard Insights")

    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Across 10 sessions, EvalTutor demonstrates a clear and consistent improvement in AI tutoring quality — with Quality Scores rising from 62 → 91 and Hallucination Risk dropping from 38% → 8% as students engage more deeply with the platform."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = TEXT_MUTED

    img_res = get_media_path("image-9-1.png")
    if img_res:
        slide.shapes.add_picture(img_res, Inches(0.8), Inches(1.8), width=Inches(11.733), height=Inches(4.2))

    add_stat_box(slide, Inches(0.8), Inches(6.2), Inches(3.7), Inches(1.1), "91", "Peak Quality Score", "Achieved by Session 10 on the 0-100 scale.", TEXT_DARK)
    add_stat_box(slide, Inches(4.8), Inches(6.2), Inches(3.7), Inches(1.1), "8%", "Hallucination Risk", "Reduced from 38% at Session 1 (79% drop).", TEXT_DARK)
    add_stat_box(slide, Inches(8.8), Inches(6.2), Inches(3.7), Inches(1.1), "3", "Evaluation Dimensions", "Quality Score, Hallucination Risk, Dependency.", TEXT_DARK)

    # ════════════════════════════════════════════════
    # SLIDE 10: Challenges Faced (Matches PDF Page 10)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "Challenges Faced & Solutions Implemented")

    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.6))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Four significant engineering challenges were encountered and resolved during development — each solution strengthening the platform's robustness and accuracy."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_MUTED

    # Table simulation
    t_l, t_t, t_w, t_h = add_card(slide, Inches(0.8), Inches(2.0), Inches(11.7), Inches(3.8), "", CARD_BG, BORDER_MUTED, TEXT_DARK)
    
    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.7), Inches(0.6))
    header.fill.solid()
    header.fill.fore_color.rgb = ACCENT_BLUE
    header.line.fill.background()
    
    tb_th = slide.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(0.5))
    tf_th = tb_th.text_frame
    p_th = tf_th.paragraphs[0]
    p_th.text = "Challenge                                                                    Solution Implemented"
    p_th.font.name = FONT_TITLE
    p_th.font.size = Pt(14)
    p_th.font.bold = True
    p_th.font.color.rgb = RGBColor(255,255,255)

    tb_t = slide.shapes.add_textbox(Inches(0.9), Inches(2.7), Inches(11.5), Inches(3.0))
    tf_t = tb_t.text_frame
    
    challs = [
        ("Groq model deprecation mid-development", "Migrated to Llama-3.3-70B-Versatile — faster and more accurate than the deprecated model."),
        ("Translation accuracy for technical terms", "English-first middleware layer with context-aware prompts to preserve technical terminology."),
        ("Poor AI response formatting", "Prompt engineering with structured Markdown and JSON output templates for consistent formatting."),
        ("Measuring educational quality objectively", "AI-based evaluation engine with a multi-dimensional scoring rubric covering accuracy, clarity, and depth.")
    ]
    for idx, (chall, sol) in enumerate(challs):
        p = tf_t.add_paragraph() if idx > 0 else tf_t.paragraphs[0]
        run_c = p.add_run()
        run_c.text = f"{chall:<65}"
        run_c.font.color.rgb = TEXT_MUTED
        run_c.font.bold = False
        
        run_s = p.add_run()
        run_s.text = sol
        run_s.font.color.rgb = TEXT_DARK
        run_s.font.bold = True
        
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.space_after = Pt(18)

    banner_l, banner_t, banner_w, banner_h = add_card(slide, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.8), "", CARD_BG_ALT, BORDER_MUTED, TEXT_DARK)
    tb_bn = slide.shapes.add_textbox(banner_l, banner_t, banner_w, banner_h)
    tf_bn = tb_bn.text_frame
    p_bn = tf_bn.paragraphs[0]
    p_bn.text = "💡 Each challenge encountered during development directly informed a stronger, more production-ready architecture — making EvalTutor more resilient and accurate than originally designed."
    p_bn.font.name = FONT_BODY
    p_bn.font.size = Pt(12)
    p_bn.font.color.rgb = TEXT_DARK

    # ════════════════════════════════════════════════
    # SLIDE 11: Advantages & Future Scope (Matches PDF Page 11)
    # ════════════════════════════════════════════════
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_header(slide, "Advantages & Future Scope")

    c1_left, c1_top, c1_width, c1_height = add_card(slide, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "✅ Current Advantages", ACCENT_BLUE, ACCENT_BLUE, RGBColor(255,255,255))
    tb_c1 = slide.shapes.add_textbox(c1_left, c1_top, c1_width, c1_height)
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    
    advs = [
        "Improves AI tutoring quality with real-time evaluation after every response.",
        "Detects and flags hallucinations before they mislead students.",
        "Encourages conceptual learning over answer copying.",
        "Reduces unhealthy student dependency on AI.",
        "Supports regional languages — inclusive for all learners.",
        "Provides actionable learning analytics for educators."
    ]
    for idx, adv in enumerate(advs):
        p = tf_c1.add_paragraph() if idx > 0 else tf_c1.paragraphs[0]
        p.text = "•  " + adv
        p.font.name = FONT_BODY
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(255,255,255)
        p.space_after = Pt(14)

    c2_left, c2_top, c2_width, c2_height = add_card(slide, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "🚀 Future Scope", CARD_BG, BORDER_MUTED, TEXT_DARK)
    tb_c2 = slide.shapes.add_textbox(c2_left, c2_top, c2_width, c2_height)
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    
    futs = [
        ("Voice-Based Tutoring", "Speech-to-text integration for hands-free learning."),
        ("OCR Support", "Handwritten question input via optical character recognition."),
        ("Mobile Application", "Native iOS & Android apps for on-the-go learning."),
        ("LMS Integration", "Moodle and Google Classroom compatibility."),
        ("Personalized AI Paths", "Adaptive learning journeys per student profile."),
        ("Offline Multilingual Models", "Low-connectivity region support.")
    ]
    for idx, (title, desc) in enumerate(futs):
        p = tf_c2.add_paragraph() if idx > 0 else tf_c2.paragraphs[0]
        p.text = "→ "
        run_t = p.add_run()
        run_t.text = f"{title} — "
        run_t.font.bold = True
        run_t.font.color.rgb = TEXT_DARK
        
        run_d = p.add_run()
        run_d.text = desc
        run_d.font.bold = False
        run_d.font.color.rgb = TEXT_MUTED
        
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.space_after = Pt(14)

    # Save presentation
    try:
        prs.save(output_path)
        print(f"Presentation exactly matched to PDF format saved successfully to: {output_path}")
    except PermissionError:
        fallback_path = output_path.replace(".pptx", "_Updated.pptx")
        prs.save(fallback_path)
        print(f"File was locked by another app. Saved successfully to: {fallback_path}")

if __name__ == "__main__":
    out = "EvalTutor-AI-Powered-Multilingual-Tutoring-Evaluation-Platform.pptx"
    if len(sys.argv) > 1:
        out = sys.argv[1]
    create_presentation(out)
